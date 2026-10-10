from openpyxl import load_workbook

from itgc.engine import run_audit

from .conftest import BASELINE


def test_full_audit_writes_deliverables(tmp_path):
    summary = run_audit(BASELINE, tmp_path, figures=False)
    assert "not effective: 12" in summary and "39/39" in summary
    wb = load_workbook(tmp_path / "ITGC_workpapers.xlsx")
    for sheet in ["Index", "Lead", "RCM", "Sampling", "SAD", "Evidence", "PBC", "AC-01", "OPS-04"]:
        assert sheet in wb.sheetnames
    letter = (tmp_path / "management_letter.md").read_text(encoding="utf-8")
    for section in ["**Condition.**", "**Criteria.**", "**Cause", "**Effect / risk.**", "**Recommendation.**",
                    "Significant deficiency", "DSS05.04", "AC-2"]:
        assert section in letter
