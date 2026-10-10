"""Generate the Excel workpaper file (openpyxl) in a Big-4 style layout.

Sheets: Index, Lead, RCM, Sampling, one test sheet per control, SAD (summary of
deficiencies), Evidence (hash manifest), PBC (request list).
"""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .dates import PERIOD_END, PERIOD_START
from .sampling import FREQUENCY_TABLE, detection_probability

NAVY = PatternFill("solid", fgColor="1F3864")
GREY = PatternFill("solid", fgColor="E7E6E6")
RED = PatternFill("solid", fgColor="F8CBAD")
GREEN = PatternFill("solid", fgColor="C6E0B4")
AMBER = PatternFill("solid", fgColor="FFE699")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
THIN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
WRAP = Alignment(wrap_text=True, vertical="top")
TICK = {"Pass": "✓", "Exception": "X", "N/A": "n/a"}

PBC_USE = {
    "hr_roster": "AC-01, AC-02, AC-07", "access_requests": "AC-01", "transfer_tickets": "AC-01",
    "termination_tickets": "AC-02", "authorized_privileged_access": "AC-03", "uar_signoffs": "AC-04",
    "uar_detail": "AC-04", "service_account_inventory": "AC-05, AC-07", "change_tickets": "CM-01, CM-03",
    "backup_job_log": "OPS-01", "restore_test_log": "OPS-02", "batch_job_log": "OPS-03", "incidents": "OPS-04",
}


def _header(ws, row: int, values: list[str]) -> None:
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.fill, c.font, c.alignment, c.border = NAVY, WHITE_BOLD, WRAP, THIN


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _kv(ws, row: int, key: str, value) -> int:
    ws.cell(row=row, column=1, value=key).font = BOLD
    c = ws.cell(row=row, column=2, value=value)
    c.alignment = WRAP
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
    return row + 1


def _fmt_detail(d: dict) -> str:
    parts = []
    for k, v in d.items():
        if v in (None, "", [], {}):
            continue
        if isinstance(v, (list, dict)):
            v = json.dumps(v)
        parts.append(f"{k}={v}")
    return "; ".join(parts)


def write_workpapers(path: Path, ev, rcm: dict, results: dict, evals, aggs) -> Path:
    wb = Workbook()
    controls = {c["id"]: c for c in rcm["controls"]}
    cls = {e.control_id: e.classification for e in evals}

    # ---------------------------------------------------------------- Index
    ws = wb.active
    ws.title = "Index"
    ws["A1"] = "Kestrel Valve Co. (fictional) - FY2026 IT General Controls testing"
    ws["A1"].font = Font(bold=True, size=14)
    r = 3
    for k, v in [("Audit period", f"{PERIOD_START} to {PERIOD_END}"),
                 ("Evidence bundle", ev.root.name),
                 ("Evidence collected", ev.manifest.get("collected_at")),
                 ("Prepared by", "Ali Askari (lab engagement - synthetic company)"),
                 ("Prepared on", date.today().isoformat()),
                 ("Reviewed by", "Not reviewed (portfolio lab; a real engagement requires senior/manager review)"),
                 ("Criteria", rcm["engagement"]["criteria"]),
                 ("Data note", "All people, tickets and logs are synthetic; system evidence was pulled live from a "
                               "lab Samba AD domain controller, PostgreSQL ERP database and git/CI repository.")]:
        r = _kv(ws, r, k, v)
    r += 1
    ws.cell(row=r, column=1, value="Tickmark legend").font = BOLD
    for t, meaning in [("✓", "Attribute agreed / tested without exception"), ("X", "Exception noted"),
                       ("n/a", "Attribute not applicable to this item")]:
        r += 1
        ws.cell(row=r, column=1, value=t)
        ws.cell(row=r, column=2, value=meaning)
    r += 2
    ws.cell(row=r, column=1, value="Sheets").font = BOLD
    for name, desc in [("Lead", "Lead schedule: conclusion per control"), ("RCM", "Risk & control matrix + framework mapping"),
                       ("Sampling", "Sample-size methodology and selections"), ("<control id>", "Test sheet per control"),
                       ("SAD", "Summary of deficiencies and severity evaluation"), ("Evidence", "SHA-256 evidence manifest"),
                       ("PBC", "Provided-by-client request list")]:
        r += 1
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=desc)
    _widths(ws, [22, 90])

    # ---------------------------------------------------------------- Lead
    ws = wb.create_sheet("Lead")
    cols = ["Control", "Name", "SOX ITGC domain", "Key", "Approach", "Population", "Tested",
            "Exceptions (sample/test)", "Exceptions (final)", "Conclusion", "Deficiency classification"]
    _header(ws, 1, cols)
    for i, (cid, res) in enumerate(results.items(), start=2):
        vals = [cid, res["name"], res["domain"], "Yes" if res["key"] else "No", res["approach"],
                res["population_size"], res["tested"], len(res["sample_exception_items"]),
                len(res["final_exception_items"]), res["final_conclusion"], cls.get(cid, "-")]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.border, c.alignment = THIN, WRAP
        ws.cell(row=i, column=10).fill = GREEN if res["final_conclusion"] == "Effective" else RED
        if cls.get(cid) == "Significant deficiency":
            ws.cell(row=i, column=11).fill = RED
        elif cid in cls:
            ws.cell(row=i, column=11).fill = AMBER
    _widths(ws, [9, 42, 26, 6, 10, 11, 8, 12, 11, 14, 24])
    ws.freeze_panes = "B2"

    # ---------------------------------------------------------------- RCM
    ws = wb.create_sheet("RCM")
    cols = ["Control", "Name", "Domain", "Risk(s)", "Control objective", "Control activity", "Owner", "Frequency",
            "Nature", "Type", "Key", "Risk rating", "Approach", "COBIT 2019", "NIST 800-53 r5", "SOC 2 TSC", "Criteria"]
    _header(ws, 1, cols)
    for i, c in enumerate(rcm["controls"], start=2):
        vals = [c["id"], c["name"], c["domain"], ", ".join(c["risk"]), c["objective"], c["description"], c["owner"],
                c["frequency"], c["nature"], c["type"], "Yes" if c.get("key") else "No", c["risk_rating"], c["approach"],
                ", ".join(c["cobit"]), ", ".join(c["nist"]), ", ".join(c["soc2"]), c["criteria"]]
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(row=i, column=j, value=v)
            cell.alignment, cell.border = WRAP, THIN
    _widths(ws, [9, 30, 18, 14, 45, 45, 18, 16, 14, 11, 6, 9, 10, 16, 18, 14, 30])
    ws.freeze_panes = "B2"
    r = len(rcm["controls"]) + 3
    ws.cell(row=r, column=1, value="Risks").font = BOLD
    for rid, text in rcm["risks"].items():
        r += 1
        ws.cell(row=r, column=1, value=rid)
        ws.cell(row=r, column=2, value=text)

    # ---------------------------------------------------------------- Sampling
    ws = wb.create_sheet("Sampling")
    _header(ws, 1, ["Frequency band", "Low risk n", "High risk n"])
    for i, (band, (lo, hi)) in enumerate(FREQUENCY_TABLE.items(), start=2):
        ws.cell(row=i, column=1, value=band)
        ws.cell(row=i, column=2, value=lo)
        ws.cell(row=i, column=3, value=hi)
    r = len(FREQUENCY_TABLE) + 3
    ws.cell(row=r, column=1, value="Event-driven populations use the band with a similar annual count "
                                   "(<=4 quarterly, <=12 monthly, <=52 weekly, <=365 daily). Zero expected deviations.").alignment = WRAP
    r += 2
    _header(ws, r, ["Control", "Population N", "Sample n", "Selection", "Items selected", "Exceptions in population (analytic)",
                    "P(sample finds >=1)", "Exceptions found by sample"])
    for cid, res in results.items():
        if res["approach"] != "sample" or res.get("analytic_exception_items") is None:
            continue
        r += 1
        k = len(res["analytic_exception_items"])
        p = detection_probability(res["population_size"], k, res["tested"])
        vals = [cid, res["population_size"], res["tested"], res["selection"], ", ".join(row["item"] for row in res["rows"]),
                k, round(p, 3), len(res["sample_exception_items"])]
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=j, value=v)
            cell.alignment, cell.border = WRAP, THIN
    _widths(ws, [16, 12, 10, 40, 50, 16, 14, 14])

    # ---------------------------------------------------------------- per-control sheets
    for cid, res in results.items():
        c = controls[cid]
        ws = wb.create_sheet(cid)
        ws["A1"] = f"{cid} - {c['name']}"
        ws["A1"].font = Font(bold=True, size=13)
        r = 3
        for k, v in [("Control objective", c["objective"]), ("Control activity", c["description"]),
                     ("Owner / frequency", f"{c['owner']} / {c['frequency']}"),
                     ("Nature / type / key", f"{c['nature']} / {c['type']} / {'Key' if c.get('key') else 'Non-key'}"),
                     ("Risk(s) addressed", "; ".join(f"{x}: {rcm['risks'][x]}" for x in c["risk"])),
                     ("Criteria", c["criteria"]), ("Population", f"{res['population_desc']} (N={res['population_size']})"),
                     ("Selection", f"{res['selection']} (tested {res['tested']})")]:
            r = _kv(ws, r, k, v)
        if res.get("completeness"):
            r = _kv(ws, r, "Completeness / IPE", _fmt_detail(res["completeness"]))
        r += 1
        ws.cell(row=r, column=1, value="Procedures performed").font = BOLD
        for p in c["procedures"]:
            r += 1
            status = ("Documented, not performed (lab has no interviewees; inquiry alone is never sufficient evidence)"
                      if p["type"] == "Inquiry" else "Performed")
            ws.cell(row=r, column=1, value=p["type"])
            ws.cell(row=r, column=2, value=p["step"]).alignment = WRAP
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
            ws.cell(row=r, column=7, value=status).alignment = WRAP
        r += 2
        ws.cell(row=r, column=1, value="Attributes").font = BOLD
        codes = list(c["attributes"])
        for code in codes:
            r += 1
            ws.cell(row=r, column=1, value=code)
            ws.cell(row=r, column=2, value=c["attributes"][code])
        r += 2
        _header(ws, r, ["Item", *codes, "Result", "Detail / evidence reference", "Notes"])
        for row in res["final_rows"]:
            r += 1
            ws.cell(row=r, column=1, value=row["item"]).border = THIN
            for j, code in enumerate(codes, start=2):
                v = row["attrs"].get(code, "N/A")
                cell = ws.cell(row=r, column=j, value=TICK.get(v, v))
                cell.alignment = Alignment(horizontal="center")
                cell.border = THIN
                if v == "Exception":
                    cell.fill = RED
            res_cell = ws.cell(row=r, column=len(codes) + 2, value="Exception" if row["exception"] else "No exception")
            res_cell.fill = RED if row["exception"] else GREEN
            ws.cell(row=r, column=len(codes) + 3, value=_fmt_detail(row["detail"])).alignment = WRAP
            ws.cell(row=r, column=len(codes) + 4, value=" ".join(row["notes"])).alignment = WRAP
        if not res["final_rows"]:
            r += 1
            ws.cell(row=r, column=1, value="No items met the exception criteria in the population (see notes).")
        r += 2
        for n in res["notes"]:
            ws.cell(row=r, column=1, value="Note").font = BOLD
            ws.cell(row=r, column=2, value=n).alignment = WRAP
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=10)
            r += 1
        r += 1
        concl = (f"Conclusion: {res['final_conclusion']}. "
                 + (f"Sample-only result: {res['sample_conclusion']}. " if res["approach"] == "sample" else "")
                 + (f"{len(res['final_exception_items'])} exception(s); see SAD for severity." if res["final_exception_items"]
                    else "No exceptions noted."))
        cell = ws.cell(row=r, column=1, value=concl)
        cell.font = BOLD
        cell.fill = GREEN if res["final_conclusion"] == "Effective" else RED
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
        _widths(ws, [26] + [7] * len(codes) + [14, 80, 50])

    # ---------------------------------------------------------------- SAD
    ws = wb.create_sheet("SAD")
    _header(ws, 1, ["Control", "Name", "Domain", "Key", "Exceptions / tested", "Realised", "Pervasive",
                    "Compensating control effective", "Classification", "MW consideration", "Rationale"])
    for i, e in enumerate(evals, start=2):
        vals = [e.control_id, e.name, e.domain, "Yes" if e.key else "No", f"{e.exceptions}/{e.tested}",
                "Yes" if e.realised else "No", "Yes" if e.pervasive else "No",
                ("Yes" if e.compensating_effective else "No") if e.compensating else "None designed",
                e.classification, "Escalate" if e.mw_consideration else "-", " ".join(e.rationale)]
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(row=i, column=j, value=v)
            cell.alignment, cell.border = WRAP, THIN
        ws.cell(row=i, column=9).fill = RED if e.classification == "Significant deficiency" else AMBER
    r = len(evals) + 4
    ws.cell(row=r, column=1, value="Aggregation by risk").font = BOLD
    r += 1
    _header(ws, r, ["Risk", "Description", "Controls not effective", "Aggregate classification", "Rationale"])
    for a in aggs:
        r += 1
        for j, v in enumerate([a.risk_id, a.risk, ", ".join(a.controls), a.classification, a.rationale], start=1):
            cell = ws.cell(row=r, column=j, value=v)
            cell.alignment, cell.border = WRAP, THIN
    _widths(ws, [9, 30, 18, 6, 10, 9, 9, 14, 20, 12, 90])

    # ---------------------------------------------------------------- Evidence
    ws = wb.create_sheet("Evidence")
    _header(ws, 1, ["File", "SHA-256", "Records", "Source", "Method"])
    for i, f in enumerate(ev.manifest.get("files", []), start=2):
        for j, v in enumerate([f["file"], f["sha256"], f.get("records"), f.get("source", "PBC (management)"),
                               f.get("method", "Provided by client")], start=1):
            ws.cell(row=i, column=j, value=v).alignment = WRAP
    _widths(ws, [38, 68, 9, 30, 60])

    # ---------------------------------------------------------------- PBC
    ws = wb.create_sheet("PBC")
    _header(ws, 1, ["#", "Request", "Rows received", "Used by", "Status"])
    for i, (name, rows) in enumerate(sorted(ev.pbc.items()), start=2):
        for j, v in enumerate([i - 1, name.replace("_", " "), len(rows), PBC_USE.get(name, ""), "Received"], start=1):
            ws.cell(row=i, column=j, value=v).border = THIN
    _widths(ws, [5, 34, 13, 18, 12])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
